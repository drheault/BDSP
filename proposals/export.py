"""Word (.docx) and PDF export of a proposal, including the fee table and Gantt."""

from io import BytesIO

from django.utils import timezone

from core.models import CompanyProfile
from scheduling.gantt import build_gantt


def _hex(color: str):
    color = color.lstrip("#")
    return tuple(int(color[i : i + 2], 16) for i in (0, 2, 4))


def _fee_rows(proposal):
    rows = []
    for item in proposal.items.prefetch_related("lines__employee", "lines__role"):
        rows.append(("item", item))
        for line in item.lines.all():
            rows.append(("line", line))
    return rows


# ----------------------------------------------------------------------------
# DOCX
# ----------------------------------------------------------------------------
def build_docx(proposal) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    company = CompanyProfile.get_solo()
    primary = RGBColor(*_hex(company.primary_color))
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2)
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)

    # cover
    if company.logo:
        try:
            doc.add_picture(company.logo.path, height=Cm(2.2))
        except Exception:
            pass
    p = doc.add_paragraph()
    r = p.add_run(company.name)
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = primary
    meta = [company.tagline, company.address, company.phone, company.email, company.website]
    for m in meta:
        if m:
            doc.add_paragraph(m).paragraph_format.space_after = Pt(0)

    doc.add_paragraph()
    k = doc.add_paragraph()
    kr = k.add_run(f"PROPOSAL {proposal.number} · {proposal.get_service_type_display().upper()}")
    kr.bold = True
    kr.font.size = Pt(9)
    kr.font.color.rgb = RGBColor(*_hex(company.accent_color))
    h = doc.add_heading(proposal.title, level=0)
    for run in h.runs:
        run.font.color.rgb = primary
    if proposal.summary:
        doc.add_paragraph(proposal.rendered_summary)

    t = doc.add_table(rows=0, cols=2)
    t.style = "Light List"
    pairs = [
        ("Prepared for", f"{proposal.client or '—'}" + (f"\n{proposal.contact.name}" + (f", {proposal.contact.title}" if proposal.contact.title else "") if proposal.contact else "")),
        ("Date", timezone.localdate().strftime("%d %B %Y")),
    ]
    if proposal.project_number:
        pairs.append(("Project number", proposal.project_number))
    if proposal.location:
        pairs.append(("Location", proposal.location))
    if proposal.owner:
        pairs.append(("Prepared by", proposal.owner.get_full_name() or proposal.owner.username))
    pairs.append(("Currency", proposal.cur.code))
    for a, b in pairs:
        row = t.add_row().cells
        row[0].text, row[1].text = a, b
        row[0].paragraphs[0].runs[0].bold = True

    n = 0
    for s in proposal.sections.all():
        n += 1
        doc.add_heading(f"{n}. {s.title}", level=1)
        for para in (s.rendered_body or "").split("\n"):
            doc.add_paragraph(para)

    services = list(proposal.services.all())
    if services:
        n += 1
        doc.add_heading(f"{n}. Scope of services", level=1)
        for sv in services:
            doc.add_paragraph(sv.name, style="List Bullet")

    tasks = list(proposal.tasks.all())
    if tasks:
        n += 1
        doc.add_heading(f"{n}. Schedule", level=1)
        png = render_gantt_png(tasks)
        if png:
            doc.add_picture(BytesIO(png), width=Cm(16.5))
        tt = doc.add_table(rows=1, cols=4)
        tt.style = "Light Grid"
        for i, hdr in enumerate(("Task", "Start", "End", "Phase")):
            tt.rows[0].cells[i].text = hdr
        for task in tasks:
            c = tt.add_row().cells
            c[0].text = ("◆ " if task.is_milestone else "") + task.name
            c[1].text = task.start.strftime("%d %b %Y")
            c[2].text = "" if task.is_milestone else task.end.strftime("%d %b %Y")
            c[3].text = task.get_phase_display()

    rows = _fee_rows(proposal)
    if rows or proposal.manual_value is not None:
        n += 1
        doc.add_heading(f"{n}. Fee", level=1)
        if rows:
            ft = doc.add_table(rows=1, cols=5)
            ft.style = "Light Grid"
            for i, hdr in enumerate(("Item / line", "Who / description", "Hours / qty", "Rate", "Amount")):
                ft.rows[0].cells[i].text = hdr
            for kind, obj in rows:
                c = ft.add_row().cells
                if kind == "item":
                    c[0].text = f"{obj.code} {obj.name}".strip()
                    c[4].text = proposal.money(obj.total)
                    for cell in c:
                        for p_ in cell.paragraphs:
                            for r_ in p_.runs:
                                r_.bold = True
                else:
                    c[0].text = "    " + obj.get_kind_display()
                    c[1].text = obj.label if obj.kind == "labour" else obj.description
                    c[2].text = f"{obj.hours:g}" if obj.kind == "labour" else f"{obj.quantity:g}"
                    c[3].text = proposal.money(obj.unit_cost, 2)
                    c[4].text = proposal.money(obj.amount)
                for i in (2, 3, 4):
                    c[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            summary = [("Labour", proposal.labour_total), ("Expenses & other", proposal.expense_total)]
            if proposal.contingency_percent:
                summary.append((f"Contingency ({proposal.contingency_percent:g}%)", proposal.contingency_amount))
            summary.append(("Total fee", proposal.fee_total))
            for label, amt in summary:
                c = ft.add_row().cells
                c[0].text = label
                c[4].text = proposal.money(amt)
                c[4].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for cell in (c[0], c[4]):
                    cell.paragraphs[0].runs[0].bold = True
        else:
            doc.add_paragraph(f"Total fee: {proposal.money(proposal.fee_total)}")

    doc.add_paragraph()
    f = doc.add_paragraph(company.proposal_footer)
    f.runs[0].font.size = Pt(8)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------
# PDF (ReportLab)
# ----------------------------------------------------------------------------
def gantt_drawing(tasks, scale: float = 1.0, max_width=None):
    """Build the Gantt as a ReportLab vector Drawing (usable as a PDF flowable)."""
    from reportlab.graphics.shapes import Drawing, Line, Polygon, Rect, String

    g = build_gantt(tasks, max_width=max_width)
    if not g:
        return None
    d = Drawing(g["width"] * scale, g["height"] * scale)
    d.scale(scale, scale)
    H = g["height"]

    def y(v):  # flip y axis
        return H - v

    for m in g["months"]:
        d.add(Rect(m["x"], y(22), m["w"], 22, fillColor=_rl("#f3f6f9"), strokeColor=_rl("#e2e8ee"), strokeWidth=0.5))
        d.add(String(m["x"] + 4, y(15), m["label"], fontSize=7, fillColor=_rl("#44535f")))
    for w in g["weeks"]:
        if g["day_px"] >= 2:
            d.add(Line(w["x"], y(22), w["x"], 0, strokeColor=_rl("#e2e8ee"), strokeWidth=0.5))
    for r in g["rows"]:
        d.add(Line(0, y(r["y"] + g["row_h"]), g["width"], y(r["y"] + g["row_h"]), strokeColor=_rl("#e2e8ee"), strokeWidth=0.5))
        d.add(String(8, y(r["cy"] + 3), r["task"].name[:34], fontSize=7.5, fontName="Helvetica-Bold" if r["is_milestone"] else "Helvetica", fillColor=_rl("#16212b")))
        if r["is_milestone"]:
            cx, cy = r["x"], y(r["cy"])
            d.add(Polygon([cx, cy + 6, cx + 6, cy, cx, cy - 6, cx - 6, cy], fillColor=_rl(r["color"]), strokeColor=None))
            d.add(String(cx + 9, cy - 2.5, r["label"], fontSize=6.5, fillColor=_rl("#74838f")))
        else:
            d.add(Rect(r["x"], y(r["y"] + g["row_h"] - 7), r["w"], g["row_h"] - 14, fillColor=_rl(r["color"]), strokeColor=None, rx=3, ry=3))
            if r["w"] > 92:
                d.add(String(r["x"] + 5, y(r["cy"] + 3), r["label"], fontSize=6.5, fillColor=_rl("#ffffff")))
    if g["today_x"]:
        d.add(Line(g["today_x"], y(22), g["today_x"], 0, strokeColor=_rl("#e07a1f"), strokeWidth=1, strokeDashArray=[3, 2]))
    return d


def render_gantt_png(tasks) -> bytes | None:
    """Rasterise the Gantt for Word. Needs a renderPM backend (rlPyCairo); returns None otherwise."""
    d = gantt_drawing(tasks, scale=2, max_width=640)
    if d is None:
        return None
    try:
        from reportlab.graphics import renderPM

        return renderPM.drawToString(d, fmt="PNG")
    except Exception:
        return None


def _rl(hexcolor):
    from reportlab.lib import colors

    return colors.HexColor(hexcolor)


def build_pdf(proposal) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    company = CompanyProfile.get_solo()
    primary = colors.HexColor(company.primary_color)
    accent = colors.HexColor(company.accent_color)
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontSize=9.5, leading=13.5)
    small = ParagraphStyle("small", parent=body, fontSize=8, textColor=colors.HexColor("#555555"))
    h1 = ParagraphStyle("h1", parent=ss["Title"], fontSize=20, leading=24, textColor=primary, alignment=0, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=12.5, textColor=primary, spaceBefore=12, spaceAfter=4)
    kicker = ParagraphStyle("kicker", parent=body, fontSize=8, textColor=accent, spaceAfter=2)
    right = ParagraphStyle("right", parent=small, alignment=TA_RIGHT)

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm, title=f"{proposal.number} {proposal.title}", author=company.name)
    story = []

    # cover band
    left = []
    if company.logo:
        try:
            left.append(Image(company.logo.path, width=4.5 * cm, height=1.8 * cm, kind="proportional"))
        except Exception:
            pass
    left.append(Paragraph(f"<b>{company.name}</b>", ParagraphStyle("co", parent=body, textColor=primary, fontSize=11)))
    co = "<br/>".join(x for x in (company.tagline, company.address.replace("\n", "<br/>"), company.phone, company.email, company.website) if x)
    cover = Table([[left, Paragraph(co, right)]], colWidths=[9 * cm, 8 * cm])
    cover.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, 0), 2, primary), ("BOTTOMPADDING", (0, 0), (-1, -1), 10)]))
    story += [cover, Spacer(1, 14)]
    story.append(Paragraph(f"PROPOSAL {proposal.number} · {proposal.get_service_type_display().upper()}", kicker))
    story.append(Paragraph(proposal.title, h1))
    if proposal.summary:
        story.append(Paragraph(proposal.rendered_summary.replace("\n", "<br/>"), body))
    meta = [["Prepared for", f"{proposal.client or '—'}" + (f"<br/>{proposal.contact.name}" if proposal.contact else "")], ["Date", timezone.localdate().strftime("%d %B %Y")]]
    if proposal.project_number:
        meta.append(["Project number", proposal.project_number])
    if proposal.location:
        meta.append(["Location", proposal.location])
    if proposal.owner:
        meta.append(["Prepared by", proposal.owner.get_full_name() or proposal.owner.username])
    meta.append(["Currency", proposal.cur.code])
    mt = Table([[Paragraph(f"<b>{a}</b>", small), Paragraph(b, body)] for a, b in meta], colWidths=[3.5 * cm, 13.5 * cm])
    mt.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("TOPPADDING", (0, 0), (-1, -1), 2)]))
    story += [Spacer(1, 8), mt]

    n = 0
    for s in proposal.sections.all():
        n += 1
        story.append(Paragraph(f"{n}. {s.title}", h2))
        for para in (s.rendered_body or "").split("\n"):
            story.append(Paragraph(para.replace("&", "&amp;").replace("<", "&lt;") or "&nbsp;", body))

    services = list(proposal.services.all())
    if services:
        n += 1
        story.append(Paragraph(f"{n}. Scope of services", h2))
        for sv in services:
            story.append(Paragraph(f"• {sv.name}", body))

    tasks = list(proposal.tasks.all())
    if tasks:
        n += 1
        story.append(Paragraph(f"{n}. Schedule", h2))
        drawing = gantt_drawing(tasks, max_width=17 * cm)
        if drawing is not None:
            story.append(drawing)

    rows = _fee_rows(proposal)
    if rows or proposal.manual_value is not None:
        n += 1
        block = [Paragraph(f"{n}. Fee", h2)]
        if rows:
            data = [["Item / line", "Who / description", "Hours / qty", "Rate", "Amount"]]
            styles = [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f6f9")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#44535f")),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
                ("LINEBELOW", (0, 0), (-1, 0), 0.75, colors.HexColor("#cdd6de")),
                ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#e2e8ee")),
                ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
            for kind, obj in rows:
                i = len(data)
                if kind == "item":
                    data.append([f"{obj.code} {obj.name}".strip(), "", "", "", proposal.money(obj.total)])
                    styles += [("FONTNAME", (0, i), (-1, i), "Helvetica-Bold"), ("BACKGROUND", (0, i), (-1, i), colors.HexColor("#fafbfc"))]
                else:
                    data.append([
                        "    " + obj.get_kind_display(),
                        obj.label if obj.kind == "labour" else obj.description,
                        f"{obj.hours:g}" if obj.kind == "labour" else f"{obj.quantity:g}",
                        proposal.money(obj.unit_cost, 2),
                        proposal.money(obj.amount),
                    ])
            summary = [("Labour", proposal.labour_total), ("Expenses & other", proposal.expense_total)]
            if proposal.contingency_percent:
                summary.append((f"Contingency ({proposal.contingency_percent:g}%)", proposal.contingency_amount))
            summary.append(("Total fee", proposal.fee_total))
            for label, amt in summary:
                i = len(data)
                data.append([label, "", "", "", proposal.money(amt)])
                styles.append(("FONTNAME", (0, i), (-1, i), "Helvetica-Bold"))
            styles.append(("LINEABOVE", (0, len(data) - len(summary)), (-1, len(data) - len(summary)), 0.75, colors.HexColor("#cdd6de")))
            ft = Table(data, colWidths=[5.2 * cm, 5.3 * cm, 2 * cm, 2.2 * cm, 2.6 * cm], repeatRows=1)
            ft.setStyle(TableStyle(styles))
            block.append(ft)
        else:
            block.append(Paragraph(f"Total fee: <b>{proposal.money(proposal.fee_total)}</b>", body))
        story.append(KeepTogether(block))

    story += [Spacer(1, 18), Paragraph(company.proposal_footer, small)]

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawString(2 * cm, 1.1 * cm, f"{company.name} · {proposal.number}")
        canvas.drawRightString(A4[0] - 2 * cm, 1.1 * cm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
