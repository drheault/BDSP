"""CSV / Office-Timesheets import into ActualEntry.

Expected columns (header names are matched case-insensitively, extra columns
are ignored): date, employee, hours, and optionally project_number, task,
cost, description, id. Rows are matched to proposals by project_number
(preferred) or the proposal supplied by the caller.
"""

import csv
import io
from datetime import datetime
from decimal import Decimal, InvalidOperation

from proposals.models import Employee, FeeItem, Proposal

from .models import ActualEntry

DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y/%m/%d")
ALIASES = {
    "date": {"date", "day", "work date", "workdate"},
    "employee": {"employee", "name", "resource", "person", "staff", "employee name"},
    "hours": {"hours", "hrs", "time", "duration"},
    "project_number": {"project_number", "project", "project no", "project no.", "job", "job number", "project code"},
    "task": {"task", "item", "wbs", "activity", "phase", "deliverable"},
    "cost": {"cost", "amount", "value"},
    "description": {"description", "notes", "comment", "comments", "memo"},
    "id": {"id", "entry id", "row id", "external_id", "uid"},
}


def _parse_date(s):
    s = (s or "").strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _dec(s):
    try:
        return Decimal(str(s).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None


def _normalise(headers, overrides=None):
    """Map source headers to canonical names. `overrides` is {canonical: source header}
    from the connector settings and wins over the built-in aliases."""
    out = {}
    overrides = {k: (v or "").strip().lower() for k, v in (overrides or {}).items() if v}
    for h in headers:
        key = (h or "").strip().lower()
        for canon, src in overrides.items():
            if key == src:
                out[h] = canon
        if h in out:
            continue
        for canon, names in ALIASES.items():
            if key in names:
                out[h] = canon
    return out


def _rows_from_xlsx(raw: bytes):
    """Read the first sheet of an Excel export; the first non-empty row is the header."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    headers = None
    for row in rows:
        if any(c not in (None, "") for c in row):
            headers = [str(c).strip() if c is not None else "" for c in row]
            break
    if not headers:
        return [], []
    out = []
    for row in rows:
        if not any(c not in (None, "") for c in row):
            continue
        rec = {}
        for h, c in zip(headers, row):
            if not h:
                continue
            if hasattr(c, "strftime"):
                c = c.strftime("%Y-%m-%d")
            rec[h] = "" if c is None else str(c)
        out.append(rec)
    return headers, out


def read_rows(file_obj, filename: str = ""):
    """Return (headers, rows) from a CSV or Excel file-like object."""
    raw = file_obj.read()
    name = (filename or getattr(file_obj, "name", "") or "").lower()
    if name.endswith((".xlsx", ".xlsm")) or (isinstance(raw, bytes) and raw[:2] == b"PK"):
        return _rows_from_xlsx(raw)
    text = raw.decode("utf-8-sig") if isinstance(raw, bytes) else raw
    reader = csv.DictReader(io.StringIO(text))
    return reader.fieldnames or [], list(reader)


def import_csv(file_obj, source="csv", default_proposal=None, column_map=None, filename=""):
    """Import a CSV or Excel timesheet export (name kept for backwards compatibility)."""
    headers, raw_rows = read_rows(file_obj, filename)
    mapping = _normalise(headers, column_map)
    created, skipped, errors = 0, 0, []
    employees = {e.name.lower(): e for e in Employee.objects.all()}
    ext_ids = {e.initials.lower(): e for e in Employee.objects.exclude(initials="")}
    ext_ids.update({e.external_id.lower(): e for e in Employee.objects.exclude(external_id="")})

    for i, raw_row in enumerate(raw_rows, start=2):
        row = {mapping.get(k, k): (v or "").strip() for k, v in raw_row.items() if k is not None}
        date = _parse_date(row.get("date"))
        hours = _dec(row.get("hours"))
        if not date or hours is None:
            errors.append(f"Row {i}: missing/invalid date or hours")
            continue
        proposal = default_proposal
        pn = row.get("project_number")
        if pn:
            proposal = Proposal.objects.filter(project_number__iexact=pn).first() or Proposal.objects.filter(number__iexact=pn).first() or proposal
        if proposal is None:
            errors.append(f"Row {i}: no proposal for project number '{pn}'")
            continue
        ext = row.get("id", "")
        if ext and ActualEntry.objects.filter(source=source, external_id=ext).exists():
            skipped += 1
            continue
        name = row.get("employee", "")
        emp = employees.get(name.lower()) or ext_ids.get(name.lower())
        item = None
        task = row.get("task", "")
        if task:
            item = FeeItem.objects.filter(proposal=proposal).filter(name__iexact=task).first() or FeeItem.objects.filter(proposal=proposal, code__iexact=task).first()
        ActualEntry.objects.create(
            proposal=proposal, item=item, employee=emp, employee_name=name, date=date, hours=hours,
            cost=_dec(row.get("cost")) if row.get("cost") else None, description=row.get("description", "")[:255],
            source=source, external_id=ext,
        )
        created += 1
    return {"created": created, "skipped": skipped, "errors": errors}
